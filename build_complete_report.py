"""Build Complete AI Interview Coach Academic Project Report (.docx, .html, companion markdown).
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

REPORT_DIR = r"D:\code_placed\ai_interview_coach\report_output"
DIAGRAMS_DIR = os.path.join(REPORT_DIR, "diagrams")
SCREENSHOTS_DIR = os.path.join(REPORT_DIR, "screenshots")
DOCX_OUT = os.path.join(REPORT_DIR, "AI_Interview_Coach_Project_Report.docx")
HTML_OUT = os.path.join(REPORT_DIR, "AI_Interview_Coach_Project_Report.html")
COMPANION_OUT = os.path.join(REPORT_DIR, "Companion_Appendix_and_Technical_Inventory.md")

os.makedirs(REPORT_DIR, exist_ok=True)

doc = Document()

# Page Margins
for s in doc.sections:
    s.top_margin = Inches(1.0)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(1.25) # Gutter
    s.right_margin = Inches(1.0)

C_NAVY = RGBColor(30, 58, 138)
C_DARK = RGBColor(31, 41, 55)
C_MUTED = RGBColor(71, 85, 105)

def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_title(text, level=1):
    h = doc.add_paragraph()
    style_p(h, space_before=16, space_after=6)
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.bold = True
    if level == 1:
        r.font.size = Pt(15)
        r.font.color.rgb = C_NAVY
    elif level == 2:
        r.font.size = Pt(13)
        r.font.color.rgb = C_NAVY
    else:
        r.font.size = Pt(11.5)
        r.font.color.rgb = C_DARK
    return h

def add_body(text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    style_p(p, space_before=0, space_after=6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = C_DARK
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.italic = italic
    r.font.color.rgb = C_DARK
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    style_p(p, space_before=0, space_after=3)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = C_DARK
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = C_DARK
    return p

def add_fig(filename, caption_text, width_in=5.8):
    path = os.path.join(DIAGRAMS_DIR, filename) if not os.path.isabs(filename) else filename
    if os.path.exists(path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p_img, space_before=10, space_after=4)
        run_img = p_img.add_run()
        run_img.add_picture(path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p_cap, space_before=2, space_after=12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.bold = True
        r_cap.font.color.rgb = C_NAVY

def add_shot(filename, caption_text, width_in=5.8):
    path = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p_img, space_before=10, space_after=4)
        run_img = p_img.add_run()
        run_img.add_picture(path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p_cap, space_before=2, space_after=12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.bold = True
        r_cap.font.color.rgb = C_NAVY

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>')
    tcPr.append(shd)

def add_tbl(headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p, space_before=4, space_after=4)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            p = row_cells[c_idx].paragraphs[0]
            style_p(p, space_before=3, space_after=3)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = C_DARK

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_after = doc.add_paragraph()
    style_p(p_after, space_before=0, space_after=8)
    return table

def add_code(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    style_p(p, space_before=4, space_after=6, line_spacing=1.0)
    r = p.add_run(code_text)
    r.font.name = "Courier New"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="2563EB"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)

print("Report generator framework loaded.")
